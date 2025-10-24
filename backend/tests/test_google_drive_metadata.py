from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

pytest.importorskip("docx")
from docx import Document  # type: ignore

BACKEND_ROOT = Path(__file__).resolve().parents[1]
if str(BACKEND_ROOT) not in sys.path:
    sys.path.insert(0, str(BACKEND_ROOT))

from app.services.google_drive.metadata import (  # noqa: E402
    build_project_folder_name,
    extract_project_metadata,
)


def _build_metadata_doc() -> bytes:
    document = Document()
    table = document.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "시험신청번호"
    table.cell(0, 1).text = "GS-B-12-3456"
    table.cell(1, 0).text = "제조자"
    table.cell(1, 1).text = "테스트 기업"
    table.cell(2, 0).text = "제품명 및 버전"
    table.cell(2, 1).text = "제품 : 보안 플랫폼 v1.2"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_project_metadata_reads_tables() -> None:
    metadata = extract_project_metadata(_build_metadata_doc())
    assert metadata == {
        "exam_number": "GS-B-12-3456",
        "company_name": "테스트 기업",
        "product_name": "보안 플랫폼 v1.2",
    }


def test_build_project_folder_name_formats_expected_string() -> None:
    folder_name = build_project_folder_name(
        {
            "exam_number": "GS-B-12-3456",
            "company_name": "테스트 기업",
            "product_name": "보안 플랫폼",
        }
    )
    assert folder_name == "[GS-B-12-3456] 테스트 기업 - 보안 플랫폼"

